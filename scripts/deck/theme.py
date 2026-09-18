"""Shared theme and slide layouts for the workshop decks.

Dark navy base with a green accent, because the subject is forestry and a deck
about canopy condition rendered in corporate blue looks like it belongs to a
different project.

Every layout function returns the slide, so a builder can add one-off shapes
afterwards without the helper having to anticipate every case.
"""

from __future__ import annotations

import os
import time
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from icons import icon_path

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
BG = RGBColor(0x0D, 0x17, 0x14)        # deep forest navy
BG_ALT = RGBColor(0x10, 0x1F, 0x1B)    # section divider background
CARD = RGBColor(0x15, 0x26, 0x21)      # card fill
CARD_EDGE = RGBColor(0x1F, 0x35, 0x2E)  # card border
GREEN = RGBColor(0x3F, 0xC1, 0x7E)     # primary accent
BLUE = RGBColor(0x38, 0x9F, 0xF7)      # secondary accent, platform topics
AMBER = RGBColor(0xF5, 0x9E, 0x0B)     # warning and highlight
COPPER = RGBColor(0xC7, 0x7A, 0x2B)    # hardwood, harvest
RED = RGBColor(0xEF, 0x44, 0x44)       # failure
WHITE = RGBColor(0xF1, 0xF5, 0xF3)     # primary text
MUTED = RGBColor(0x9C, 0xAF, 0xA8)     # secondary text
DIM = RGBColor(0x64, 0x7C, 0x74)       # tertiary text

BRONZE = RGBColor(0xC0, 0x84, 0x4A)
SILVER = RGBColor(0xB6, 0xC2, 0xC9)
GOLD = RGBColor(0xE3, 0xB3, 0x41)

FONT = "Segoe UI"
MONO = "Consolas"

SLIDE_W = 13.333
SLIDE_H = 7.5
MARGIN = 0.9

OUT_DIR = Path(__file__).resolve().parent / "out"


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------


def new_deck() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    return prs


def blank(prs: Presentation, background: RGBColor = BG):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = background
    return slide


def text_box(slide, left, top, width, height, text, size=14, colour=WHITE,
             bold=False, align=PP_ALIGN.LEFT, font=FONT, italic=False, spacing=1.0):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0

    for index, line in enumerate(str(text).split("\n")):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.text = line
        para.alignment = align
        para.line_spacing = spacing
        run_font = para.font
        run_font.size = Pt(size)
        run_font.color.rgb = colour
        run_font.bold = bold
        run_font.italic = italic
        run_font.name = font
    return box


def bullet_list(slide, left, top, width, height, items, size=15, colour=WHITE,
                marker="\u25b8", gap=16, marker_colour=None):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0

    for index, item in enumerate(items):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.space_before = Pt(0 if index == 0 else gap)
        para.line_spacing = 1.15

        head = para.add_run()
        head.text = f"{marker}  " if marker else ""
        head.font.size = Pt(size)
        head.font.color.rgb = marker_colour or GREEN
        head.font.name = FONT
        head.font.bold = True

        body = para.add_run()
        body.text = str(item)
        body.font.size = Pt(size)
        body.font.color.rgb = colour
        body.font.name = FONT
    return box


def card(slide, left, top, width, height, fill=CARD, edge=CARD_EDGE, line_pt=1.0):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.adjustments[0] = 0.06
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = edge
    shape.line.width = Pt(line_pt)
    shape.shadow.inherit = False
    return shape


def bar(slide, left, top, width=1.15, height=0.06, colour=GREEN):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = colour
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def glyph(slide, name, left, top, size=0.44, colour=GREEN):
    return slide.shapes.add_picture(
        icon_path(name, (colour[0], colour[1], colour[2])),
        Inches(left), Inches(top), height=Inches(size), width=Inches(size),
    )


def badge(slide, left, top, diameter, text, fill=GREEN, text_colour=BG, size=15):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(diameter), Inches(diameter)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    frame = shape.text_frame
    frame.word_wrap = False
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.text = str(text)
    para.alignment = PP_ALIGN.CENTER
    para.font.size = Pt(size)
    para.font.bold = True
    para.font.color.rgb = text_colour
    para.font.name = FONT
    return shape


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()


def chrome(slide, kicker=None, page=None, total=None, icon=None, accent=GREEN):
    """Standard slide furniture: accent bar, kicker, optional icon and page number."""
    bar(slide, MARGIN, 0.62, colour=accent)
    if kicker:
        text_box(slide, MARGIN, 0.32, 8.0, 0.28, kicker.upper(), 10.5, DIM, bold=True)
    if icon:
        glyph(slide, icon, SLIDE_W - MARGIN - 0.5, 0.34, 0.5, accent)
    if page is not None:
        label = f"{page} / {total}" if total else str(page)
        text_box(slide, SLIDE_W - MARGIN - 1.6, SLIDE_H - 0.55, 1.6, 0.3, label,
                 10, DIM, align=PP_ALIGN.RIGHT)


# ---------------------------------------------------------------------------
# Slide layouts
# ---------------------------------------------------------------------------


def title_slide(prs, title, tagline, subtitle, footer, icon_row=None):
    slide = blank(prs, BG)

    # A band of domain icons instead of stock photography.
    for index, name in enumerate(icon_row or ["satellite", "layers", "lakehouse", "brain", "chart"]):
        glyph(slide, name, MARGIN + index * 0.78, 1.25, 0.52, GREEN if index % 2 == 0 else BLUE)

    bar(slide, MARGIN, 2.35, 1.6, 0.075)
    text_box(slide, MARGIN, 2.62, 11.2, 1.5, title, 44, WHITE, bold=True, spacing=1.05)
    text_box(slide, MARGIN, 4.18, 11.2, 0.6, tagline, 21, GREEN)
    text_box(slide, MARGIN, 4.95, 11.2, 0.9, subtitle, 14, MUTED, spacing=1.25)
    text_box(slide, MARGIN, SLIDE_H - 0.85, 11.2, 0.4, footer, 11, DIM)
    return slide


def section_slide(prs, number, title, subtitle, icon="layers", accent=GREEN, duration=None):
    slide = blank(prs, BG_ALT)
    glyph(slide, icon, MARGIN, 2.35, 0.85, accent)
    bar(slide, MARGIN, 3.45, 1.4, 0.075, accent)
    text_box(slide, MARGIN, 1.85, 3.0, 0.4, f"BLOCK {number}", 11, DIM, bold=True)
    text_box(slide, MARGIN, 3.72, 10.5, 1.0, title, 34, WHITE, bold=True)
    text_box(slide, MARGIN, 4.85, 10.5, 0.9, subtitle, 15, MUTED, spacing=1.25)
    if duration:
        text_box(slide, SLIDE_W - MARGIN - 3.0, 1.85, 3.0, 0.4, duration, 12, accent,
                 bold=True, align=PP_ALIGN.RIGHT)
    return slide


def bullets_slide(prs, title, items, kicker=None, icon=None, footnote=None,
                  accent=GREEN, page=None, total=None, size=15):
    slide = blank(prs)
    chrome(slide, kicker, page, total, icon, accent)
    text_box(slide, MARGIN, 0.84, 11.3, 0.7, title, 29, WHITE, bold=True)
    bullet_list(slide, MARGIN, 1.95, 11.3, 4.2, items, size, WHITE, marker_colour=accent)
    if footnote:
        card(slide, MARGIN, SLIDE_H - 1.55, 11.53, 0.85)
        text_box(slide, MARGIN + 0.35, SLIDE_H - 1.33, 10.9, 0.6, footnote, 13.5, accent, spacing=1.2)
    return slide


def cards_slide(prs, title, cards, kicker=None, icon=None, columns=3,
                accent=GREEN, page=None, total=None, subtitle=None):
    slide = blank(prs)
    chrome(slide, kicker, page, total, icon, accent)
    text_box(slide, MARGIN, 0.84, 11.3, 0.7, title, 29, WHITE, bold=True)
    top = 1.95
    if subtitle:
        text_box(slide, MARGIN, 1.62, 11.3, 0.4, subtitle, 13.5, MUTED)
        top = 2.25

    rows = (len(cards) + columns - 1) // columns
    usable = SLIDE_W - 2 * MARGIN
    gap = 0.28
    width = (usable - gap * (columns - 1)) / columns
    height = min(2.15, (SLIDE_H - top - 0.7 - gap * (rows - 1)) / rows)

    for index, entry in enumerate(cards):
        icon_name, head, body = entry
        row, col = divmod(index, columns)
        left = MARGIN + col * (width + gap)
        y = top + row * (height + gap)

        card(slide, left, y, width, height)
        if icon_name:
            glyph(slide, icon_name, left + 0.3, y + 0.28, 0.42, accent)
            head_y = y + 0.85
        else:
            head_y = y + 0.32
        text_box(slide, left + 0.3, head_y, width - 0.6, 0.4, head, 15, accent, bold=True)
        text_box(slide, left + 0.3, head_y + 0.42, width - 0.6, height - (head_y - y) - 0.55,
                 body, 12, MUTED, spacing=1.2)
    return slide


def compare_slide(prs, title, left_head, left_items, right_head, right_items,
                  kicker=None, icon=None, page=None, total=None,
                  left_accent=GREEN, right_accent=COPPER, footnote=None):
    slide = blank(prs)
    chrome(slide, kicker, page, total, icon)
    text_box(slide, MARGIN, 0.84, 11.3, 0.7, title, 29, WHITE, bold=True)

    width = (SLIDE_W - 2 * MARGIN - 0.4) / 2
    height = 3.55 if footnote else 4.35

    for offset, (head, items, accent, marker) in enumerate([
        (left_head, left_items, left_accent, "\u2713"),
        (right_head, right_items, right_accent, "\u2717"),
    ]):
        left = MARGIN + offset * (width + 0.4)
        card(slide, left, 1.95, width, height)
        bar(slide, left + 0.32, 2.24, 0.7, 0.05, accent)
        text_box(slide, left + 0.32, 2.42, width - 0.64, 0.45, head, 17, accent, bold=True)
        bullet_list(slide, left + 0.32, 3.05, width - 0.64, height - 1.25, items,
                    13.5, WHITE, marker=marker, gap=13, marker_colour=accent)

    if footnote:
        card(slide, MARGIN, SLIDE_H - 1.35, 11.53, 0.78)
        text_box(slide, MARGIN + 0.35, SLIDE_H - 1.14, 10.9, 0.5, footnote, 13.5, AMBER)
    return slide


def table_slide(prs, title, headers, rows, widths=None, kicker=None, icon=None,
                accent=GREEN, page=None, total=None, footnote=None, size=12.5):
    slide = blank(prs)
    chrome(slide, kicker, page, total, icon, accent)
    text_box(slide, MARGIN, 0.84, 11.3, 0.7, title, 29, WHITE, bold=True)

    usable = SLIDE_W - 2 * MARGIN
    widths = widths or [1 / len(headers)] * len(headers)
    widths = [w / sum(widths) * usable for w in widths]

    header_y = 2.0
    text_y = header_y + 0.42
    available = SLIDE_H - text_y - (1.5 if footnote else 0.75)
    row_h = min(0.86, available / max(len(rows), 1))

    x = MARGIN
    for width, header in zip(widths, headers):
        text_box(slide, x, header_y, width - 0.15, 0.34, header.upper(), 10.5, accent, bold=True)
        x += width
    bar(slide, MARGIN, header_y + 0.36, usable, 0.02, CARD_EDGE)

    for index, row in enumerate(rows):
        y = text_y + index * row_h
        if index % 2 == 0:
            band = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(MARGIN - 0.12), Inches(y - 0.05),
                Inches(usable + 0.24), Inches(row_h - 0.04)
            )
            band.fill.solid()
            band.fill.fore_color.rgb = CARD
            band.line.fill.background()
            band.shadow.inherit = False

        x = MARGIN
        for col, (width, value) in enumerate(zip(widths, row)):
            colour = WHITE if col == 0 else MUTED
            text_box(slide, x, y + 0.05, width - 0.15, row_h - 0.1, value, size, colour,
                     bold=(col == 0))
            x += width

    if footnote:
        text_box(slide, MARGIN, SLIDE_H - 1.05, 11.3, 0.5, footnote, 13, accent)
    return slide


def flow_slide(prs, title, steps, kicker=None, icon="pipeline", accent=GREEN,
               page=None, total=None, footnote=None):
    """Horizontal flow of labelled boxes with arrows between them."""
    slide = blank(prs)
    chrome(slide, kicker, page, total, icon, accent)
    text_box(slide, MARGIN, 0.84, 11.3, 0.7, title, 29, WHITE, bold=True)

    count = len(steps)
    usable = SLIDE_W - 2 * MARGIN
    gap = 0.42
    width = (usable - gap * (count - 1)) / count
    top, height = 2.35, 2.5

    for index, (step_icon, head, body, colour) in enumerate(steps):
        left = MARGIN + index * (width + gap)
        card(slide, left, top, width, height, edge=colour, line_pt=1.5)
        glyph(slide, step_icon, left + 0.28, top + 0.28, 0.44, colour)
        text_box(slide, left + 0.28, top + 0.88, width - 0.56, 0.4, head, 15, colour, bold=True)
        text_box(slide, left + 0.28, top + 1.32, width - 0.56, height - 1.5, body, 11.5, MUTED,
                 spacing=1.2)

        if index < count - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW,
                Inches(left + width + 0.06), Inches(top + height / 2 - 0.11),
                Inches(gap - 0.12), Inches(0.22),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = CARD_EDGE
            arrow.line.fill.background()
            arrow.shadow.inherit = False

    if footnote:
        card(slide, MARGIN, SLIDE_H - 1.5, 11.53, 0.82)
        text_box(slide, MARGIN + 0.35, SLIDE_H - 1.29, 10.9, 0.55, footnote, 13.5, accent)
    return slide


def steps_slide(prs, title, steps, kicker=None, icon="checklist", accent=GREEN,
                page=None, total=None):
    """Numbered vertical list with an optional right-aligned value per row."""
    slide = blank(prs)
    chrome(slide, kicker, page, total, icon, accent)
    text_box(slide, MARGIN, 0.84, 11.3, 0.7, title, 29, WHITE, bold=True)

    available = SLIDE_H - 1.95 - 0.75
    height = min(1.02, available / max(len(steps), 1))
    # Centre a short list rather than stranding it against the title.
    top = 1.95 + max(0.0, (available - height * len(steps)) / 2)

    for index, entry in enumerate(steps):
        head, body, value = (list(entry) + [None])[:3]
        y = top + index * height
        badge(slide, MARGIN, y + 0.04, 0.44, index + 1, accent, BG, 14)
        text_box(slide, MARGIN + 0.66, y + 0.02, 6.4, 0.34, head, 15.5, WHITE, bold=True)
        text_box(slide, MARGIN + 0.66, y + 0.36, 8.6, height - 0.4, body, 12.5, MUTED, spacing=1.2)
        if value:
            text_box(slide, SLIDE_W - MARGIN - 2.2, y + 0.06, 2.2, 0.4, value, 13, accent,
                     bold=True, align=PP_ALIGN.RIGHT)
    return slide


def code_slide(prs, title, code, caption=None, kicker=None, icon="code",
               accent=GREEN, page=None, total=None, size=None):
    slide = blank(prs)
    chrome(slide, kicker, page, total, icon, accent)
    text_box(slide, MARGIN, 0.84, 11.3, 0.7, title, 29, WHITE, bold=True)

    lines = code.strip("\n").split("\n")
    # Shrink long listings rather than letting them run off the slide.
    if size is None:
        size = 12.5 if len(lines) <= 13 else (11.0 if len(lines) <= 17 else 9.8)

    # Rendered line height is roughly 1.2x the point size, then the 1.35 line
    # spacing set below. Measuring by point size alone under-counts and clips
    # the last two lines, which is how a code card ends up cut off in the room.
    line_h = size * 1.2 * 1.35 / 72
    pad = 0.21
    height = 2 * pad + len(lines) * line_h
    card(slide, MARGIN, 1.95, 11.53, height, fill=RGBColor(0x0A, 0x12, 0x10))
    text_box(slide, MARGIN + 0.35, 1.95 + pad, 10.9, height - 2 * pad, "\n".join(lines), size, WHITE,
             font=MONO, spacing=1.35)

    if caption:
        text_box(slide, MARGIN, 1.95 + height + 0.28, 11.3, 0.9, caption, 13.5, MUTED, spacing=1.25)
    return slide


def quote_slide(prs, text, attribution=None, accent=GREEN, icon="lightbulb",
                page=None, total=None):
    slide = blank(prs, BG_ALT)
    glyph(slide, icon, MARGIN, 2.15, 0.7, accent)
    bar(slide, MARGIN, 3.1, 1.4, 0.075, accent)
    text_box(slide, MARGIN, 3.42, 11.2, 2.2, text, 26, WHITE, bold=True, spacing=1.2)
    if attribution:
        text_box(slide, MARGIN, 5.75, 11.2, 0.4, attribution, 13, MUTED)
    if page is not None:
        text_box(slide, SLIDE_W - MARGIN - 1.6, SLIDE_H - 0.55, 1.6, 0.3,
                 f"{page} / {total}" if total else str(page), 10, DIM, align=PP_ALIGN.RIGHT)
    return slide


def medallion_slide(prs, title, bronze_items, silver_items, gold_items,
                    kicker=None, page=None, total=None, footnote=None):
    """The three-layer diagram, used in both decks."""
    slide = blank(prs)
    chrome(slide, kicker, page, total, "layers", GREEN)
    text_box(slide, MARGIN, 0.84, 11.3, 0.7, title, 29, WHITE, bold=True)

    layers = [
        ("download", "BRONZE", "fidelity to source", bronze_items, BRONZE),
        ("filter", "SILVER", "comparability", silver_items, SILVER),
        ("target", "GOLD", "meaning for a decision", gold_items, GOLD),
    ]
    top, height, gap = 1.86, 1.50, 0.17
    longest = max(len(items) for _, _, _, items, _ in layers)
    size = 12.5 if longest <= 3 else 11.5
    spacing = 9 if longest <= 3 else 5

    for index, (icon_name, name, promise, items, colour) in enumerate(layers):
        y = top + index * (height + gap)
        card(slide, MARGIN, y, 11.53, height, edge=colour, line_pt=1.5)
        glyph(slide, icon_name, MARGIN + 0.32, y + 0.30, 0.44, colour)
        text_box(slide, MARGIN + 0.95, y + 0.26, 2.6, 0.36, name, 17, colour, bold=True)
        text_box(slide, MARGIN + 0.95, y + 0.66, 2.8, 0.4, promise, 11.5, DIM)
        bullet_list(slide, MARGIN + 3.9, y + 0.24, 7.3, height - 0.4, items, size, MUTED,
                    marker="\u00b7", gap=spacing, marker_colour=colour)

    if footnote:
        text_box(slide, MARGIN, SLIDE_H - 0.72, 9.4, 0.45, footnote, 12.5, GREEN)
    return slide


def closing_slide(prs, title, items, footer, icon="rocket", accent=GREEN):
    slide = blank(prs, BG_ALT)
    glyph(slide, icon, MARGIN, 1.55, 0.72, accent)
    bar(slide, MARGIN, 2.5, 1.4, 0.075, accent)
    text_box(slide, MARGIN, 2.8, 11.2, 0.8, title, 32, WHITE, bold=True)
    bullet_list(slide, MARGIN, 3.85, 11.2, 2.4, items, 15, WHITE, marker_colour=accent)
    text_box(slide, MARGIN, SLIDE_H - 0.85, 11.2, 0.4, footer, 11, DIM)
    return slide


# ---------------------------------------------------------------------------
# Save
# ---------------------------------------------------------------------------


def save(prs: Presentation, filename: str) -> Path:
    """Save, versioning the filename if the target is open in PowerPoint."""
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    path = OUT_DIR / filename
    if path.exists():
        try:
            with open(path, "ab"):
                pass
        except PermissionError:
            stamp = time.strftime("%Y%m%d-%H%M%S")
            path = OUT_DIR / f"{path.stem}-{stamp}{path.suffix}"
    prs.save(path)
    return path


def summarise(prs: Presentation, path: Path) -> None:
    with_notes = sum(
        1 for s in prs.slides
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip()
    )
    size_kb = os.path.getsize(path) / 1024
    print(f"  {path.name}")
    print(f"    {len(prs.slides)} slides, {with_notes} with speaker notes, {size_kb:.0f} KB")
