"""Build the screenshot-led manual Fabric workshop deck without rebuilding notebooks.

Usage:
    python scripts/deck/build_manual_deck.py --draft
    python scripts/deck/build_manual_deck.py
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import logging
import re
import sys
from pathlib import Path
from zipfile import ZipFile

from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from theme import BG, GREEN, MUTED, WHITE, badge, blank, new_deck, notes, text_box

ROOT = Path(__file__).resolve().parents[2]
IMAGES = ROOT / "docs/images/training/manual"
CAPTURES = ROOT / "scripts/verification/screenshots"
INK = RGBColor(0x17, 0x35, 0x2E)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
RED = RGBColor(0xCB, 0x3A, 0x24)
LOGGER = logging.getLogger(__name__)


def capture_section(number: int) -> str:
    """Label the responsibility and workflow stage shown on a slide."""
    if number in {1, 2, 10, 11, 12}:
        return "Facilitator preparation"
    if number <= 16:
        return "Student setup"
    if number <= 24:
        return "Complete Labs 00-05"
    if number == 25:
        return "Check the Gold data"
    if number <= 32:
        return "Build the native Fabric Map"
    if number <= 41:
        return "Build and test the Fabric data agent"
    return "Final checkpoint"


def handout_notes(filename: str) -> str:
    """Include the matching handout section as editable speaker notes."""
    for relative in ("docs/16-manual-upload-labs.md", "docs/12-spark-environment.md",
                     "docs/17-manual-imagery-download.md"):
        text = (ROOT / relative).read_text(encoding="utf-8")
        for section in re.split(r"(?m)(?=^#{2,3} )", text):
            if filename in section:
                directory = "manual-download" if relative.endswith("17-manual-imagery-download.md") else "manual"
                return f"Source: {relative}\nCapture: images/training/{directory}/{filename}\n\n{section.strip()}"
    raise ValueError(f"Screenshot is not referenced in the handouts: {filename}")


def focused_picture(capture: dict) -> tuple[io.BytesIO, tuple[int, int]]:
    """Crop only surrounding context, retaining every numbered callout and target."""
    directory = ROOT / capture.get("directory", "docs/images/training/manual")
    path = directory / capture["file"]
    raw = CAPTURES / directory.name / "raw" / capture["file"]
    with Image.open(raw) as original:
        raw_width, raw_height = original.size
    with Image.open(path) as annotated:
        expected_hash = hashlib.sha256(raw.read_bytes()).hexdigest()
        if annotated.info.get("Source SHA256") != expected_hash:
            raise ValueError(f"Annotated screenshot is stale: {capture['file']}")
        bounds = []
        for focus in capture["focus"]:
            left, top, width, height = focus["box"]
            badge_x, badge_y = focus["badge"]
            bounds.extend([(left - 10, top - 10, left + width + 10, top + height + 10),
                           (badge_x - 20, badge_y - 20, badge_x + 20, badge_y + 20)])
        crop = (max(0, int(min(bound[0] for bound in bounds) - 80)),
                max(0, int(min(bound[1] for bound in bounds) - 80)),
                min(raw_width, int(max(bound[2] for bound in bounds) + 80)),
                min(raw_height, int(max(bound[3] for bound in bounds) + 65)))
        crop = tuple(capture.get("deck_crop", crop))
        if not (0 <= crop[0] < crop[2] <= raw_width and 0 <= crop[1] < crop[3] <= raw_height):
            raise ValueError(f"Deck crop is outside the raw capture: {capture['file']}")
        if any(bound[0] < crop[0] or bound[1] < crop[1]
               or bound[2] > crop[2] or bound[3] > crop[3] for bound in bounds):
            raise ValueError(f"Deck crop clips a callout: {capture['file']}")
        picture = annotated.crop(crop)
    stream = io.BytesIO()
    picture.save(stream, format="PNG")
    stream.seek(0)
    return stream, picture.size


def screenshot_slide(presentation, capture: dict, number: int, total: int) -> None:
    """Create one focused screenshot step with editable numbered directions."""
    slide = blank(presentation, PAPER)
    source_number = int(capture["file"].split("-", 1)[0])
    section = "Download and upload imagery" if capture.get("directory") else capture_section(source_number)
    text_box(slide, 0.55, 0.25, 10.8, 0.25, section.upper(),
             10, INK, bold=True)
    title = capture["title"].removeprefix("Facilitator: ")
    title = title[:1].upper() + title[1:]
    text_box(slide, 0.55, 0.64, 12.2, 0.6, title, 26 if len(title) < 75 else 23,
             INK, bold=True)
    stream, (width, height) = focused_picture(capture)
    scale = min(12.15 / width, 4.65 / height)
    picture_width, picture_height = width * scale, height * scale
    slide.shapes.add_picture(stream, Inches((13.333 - picture_width) / 2),
                             Inches(1.38 + (4.65 - picture_height) / 2),
                             width=Inches(picture_width), height=Inches(picture_height))
    gap = 0.32
    column_width = (12.15 - gap * (len(capture["focus"]) - 1)) / len(capture["focus"])
    for index, focus in enumerate(capture["focus"]):
        left = 0.59 + index * (column_width + gap)
        badge(slide, left, 6.22, 0.32, index + 1, fill=RED, text_colour=WHITE, size=11)
        text_box(slide, left + 0.44, 6.18, column_width - 0.44, 0.87,
                 focus["text"], 14, INK, spacing=1.05)
    text_box(slide, 0.55, 7.19, 9.8, 0.2, "Woodlands in Fabric | Manual workshop", 9, INK)
    text_box(slide, 11.3, 7.15, 1.45, 0.25, f"{number} / {total}", 10, INK,
             align=PP_ALIGN.RIGHT)
    notes(slide, handout_notes(capture["file"]))


def validate_deck(presentation, expected_slides: int) -> None:
    """Fail on missing notes, missing pictures, or off-slide geometry."""
    if len(presentation.slides) != expected_slides:
        raise ValueError("Deck slide count does not match the captured workflow")
    for number, slide in enumerate(presentation.slides, 1):
        if not slide.notes_slide.notes_text_frame.text.strip():
            raise ValueError(f"Slide {number} has no speaker notes")
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0
                    or shape.left + shape.width > presentation.slide_width + 2
                    or shape.top + shape.height > presentation.slide_height + 2):
                raise ValueError(f"Off-slide shape on slide {number}: {shape.name}")
        if 1 < number < expected_slides and not any(shape.shape_type == 13 for shape in slide.shapes):
            raise ValueError(f"Step slide {number} is missing its screenshot")


def build(output: Path, *, draft: bool = False) -> int:
    """Build a draft of captured steps or a release only after all live gates pass."""
    evidence = json.loads((ROOT / "scripts/verification/training-manual-evidence.json").read_text(encoding="utf-8"))
    captures = json.loads((CAPTURES / "manual/annotations.json").read_text(encoding="utf-8"))["captures"]
    download_directory = "docs/images/training/manual-download"
    download_captures = json.loads((CAPTURES / "manual-download/annotations.json").read_text(encoding="utf-8"))["captures"]
    if not draft:
        if evidence["status"] != "ready_for_classroom":
            raise ValueError("Final deck requires ready_for_classroom evidence; use --draft during rehearsal")
        expected = {entry["file"] for entry in evidence["screenshots"]}
        if {entry["file"] for entry in captures} != expected:
            raise ValueError("Final deck must contain every required screenshot")
        download_evidence = json.loads((ROOT / "scripts/verification/training-manual-download-evidence.json").read_text(encoding="utf-8"))
        if download_evidence["status"] != "verified":
            raise ValueError("Final deck needs verified manual-download rehearsal evidence")
    captures.extend({**capture, "directory": download_directory} for capture in download_captures)
    order = [1, 2, 10, 11, 12, *range(3, 10), *range(13, 43)]
    captures.sort(key=lambda entry: order.index(int(entry["file"].split("-", 1)[0])))
    presentation = new_deck()
    total = len(captures) + 2
    slide = blank(presentation, BG)
    text_box(slide, 0.75, 0.62, 11.8, 0.4, "WOODLANDS | MANUAL FABRIC WORKSHOP", 12, GREEN, bold=True)
    text_box(slide, 0.75, 1.4, 11.8, 1.3, "Woodlands in Fabric", 44, WHITE, bold=True)
    text_box(slide, 0.75, 2.75, 11.7, 0.8,
             "Student notebooks. Your lakehouse. A native Map and data agent.", 23, WHITE)
    for index, (heading, detail) in enumerate([
        ("Prepare", "Upload six notebooks, then create and attach your lakehouse."),
        ("Process", "Choose manual imagery upload or STAC; complete each lab checkpoint."),
        ("Inspect", "Build the Map and test the agent against Gold SQL results."),
    ]):
        top = 4.05 + index * 0.75
        badge(slide, 0.8, top, 0.4, index + 1, size=14)
        text_box(slide, 1.4, top - 0.03, 2.0, 0.42, heading, 21, WHITE, bold=True)
        text_box(slide, 3.4, top - 0.02, 8.9, 0.5, detail, 17, MUTED)
    label = "DRAFT: release review is not complete" if draft else "Manual imagery upload + STAC | Verified September 2026"
    text_box(slide, 0.75, 6.94, 11.8, 0.25, label, 11, GREEN)
    notes(slide, "Use the student handout alongside this deck. The first five steps are facilitator preparation. "
          "The learner sequence uploads all six notebooks before creating a lakehouse. "
          "Public imagery and synthetic stands only; no production inventory. " + label)
    for number, capture in enumerate(captures, 2):
        screenshot_slide(presentation, capture, number, total)
    slide = blank(presentation, BG)
    text_box(slide, 0.75, 0.7, 11.8, 0.7, "Your hand-in", 36, WHITE, bold=True)
    for index, instruction in enumerate([
        "Six completed notebooks with passing validation output",
        "Your lakehouse, with verified Gold tables and GeoJSON export",
        "A saved native Map with class colors, tooltips and coverage filtering",
        "A published data agent whose answers match the independent SQL checks",
    ]):
        top = 2.0 + index * 0.92
        badge(slide, 0.8, top, 0.42, index + 1, size=14)
        text_box(slide, 1.48, top - 0.03, 10.8, 0.65, instruction, 22, WHITE)
    text_box(slide, 0.8, 6.28, 11.7, 0.6,
             "Explain missing results, simulated history and offline narrative stubs. Stop your Spark session.",
             18, GREEN)
    notes(slide, "Use the hand-in checklist in the student handout. A scheduler success is not enough: "
          "inspect validations and durable data. Do not pause the shared capacity or grant external access.")
    validate_deck(presentation, total)
    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    with ZipFile(output) as archive:
        if archive.testzip() is not None:
            raise ValueError("PPTX ZIP integrity check failed")
    LOGGER.info("Built %s: %d slides, %d real screenshots, draft=%s", output, total, len(captures), draft)
    return total


def create_parser() -> argparse.ArgumentParser:
    """Create command options for draft or gated final output."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--draft", action="store_true")
    parser.add_argument("--output", type=Path)
    return parser


def main() -> int:
    """Build the new manual deck without touching the existing session decks."""
    arguments = create_parser().parse_args()
    logging.basicConfig(level=logging.INFO, format="%(message)s")
    output = arguments.output or (
        ROOT / ".copilot-tracking/build/woodlands-manual-workshop-draft.pptx" if arguments.draft
        else ROOT / "decks/woodlands-manual-workshop.pptx")
    try:
        build(output, draft=arguments.draft)
        return 0
    except (OSError, ValueError, KeyError) as error:
        LOGGER.error("Deck build failed: %s", error)
        return 1


if __name__ == "__main__":
    sys.exit(main())
